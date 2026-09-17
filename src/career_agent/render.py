"""Render a TailoredResume as Markdown, DOCX and Notion blocks.

The DOCX is deliberately plain - one column, no tables, images, text boxes or
headers/footers - because that is what applicant tracking systems parse most
reliably. Every word comes from the TailoredResume, which the ValidationGate
has already checked against the master profile.
"""

from __future__ import annotations

import io
import re
from typing import Any

from .models import MasterProfile, TailoredResume

_SAFE = re.compile(r"[^A-Za-z0-9]+")


def resume_filename(resume: TailoredResume, ext: str = "docx") -> str:
    """Company_Role_YYYY-MM-DD.<ext> (SYSTEM_SPEC 5)."""
    company = _SAFE.sub("_", resume.company).strip("_")[:40] or "Company"
    role = _SAFE.sub("_", resume.role).strip("_")[:60] or "Role"
    return f"{company}_{role}_{resume.created.isoformat()}.{ext}"


def contact_line(profile: MasterProfile) -> str:
    parts = [profile.location, profile.phone, profile.email]
    if profile.github:
        parts.append(profile.github.replace("https://", ""))
    if profile.linkedin:
        parts.append(profile.linkedin.replace("https://www.", "").replace("https://", ""))
    return "  ·  ".join(p for p in parts if p)


def to_markdown(resume: TailoredResume, profile: MasterProfile) -> str:
    lines = [f"# {profile.name.upper()}", resume.headline, contact_line(profile), ""]
    lines += ["## PROFESSIONAL SUMMARY", resume.summary, ""]
    lines.append("## SKILLS")
    for group, items in resume.skills.items():
        if items:
            lines.append(f"- **{group}:** " + " · ".join(items))
    lines.append("")
    lines.append("## PROFESSIONAL EXPERIENCE")
    for e in resume.experience:
        lines.append(f"### {e.title} — {e.employer} ({e.start} – {e.end})")
        if e.location:
            lines.append(e.location)
        lines += [f"- {b}" for b in e.bullets]
    lines.append("")
    if resume.projects:
        lines.append("## AI ENGINEERING PROJECTS")
        for p in resume.projects:
            head = f"### {p.name}"
            if p.url and p.shipped:
                head += f" — {p.url.replace('https://', '')}"
            lines.append(head)
            if p.stack:
                lines.append(p.stack)
            lines += [f"- {b}" for b in p.bullets]
        lines.append("")
    if resume.certifications:
        lines.append("## CERTIFICATIONS")
        lines += [f"- {c}" for c in resume.certifications]
        lines.append("")
    if resume.education:
        lines.append("## EDUCATION")
        lines += [f"- {e}" for e in resume.education]
    return "\n".join(lines).strip() + "\n"


def to_docx(resume: TailoredResume, profile: MasterProfile) -> bytes:
    from docx import Document  # imported lazily: optional at import time for the MCP server
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    def heading(text: str) -> None:
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = True
        run.font.size = Pt(11.5)

    name = doc.add_paragraph()
    r = name.add_run(profile.name.upper())
    r.bold = True
    r.font.size = Pt(16)
    doc.add_paragraph(resume.headline)
    doc.add_paragraph(contact_line(profile))

    heading("PROFESSIONAL SUMMARY")
    doc.add_paragraph(resume.summary)

    heading("SKILLS")
    for group, items in resume.skills.items():
        if items:
            para = doc.add_paragraph()
            para.add_run(f"{group}: ").bold = True
            para.add_run(" · ".join(items))

    heading("PROFESSIONAL EXPERIENCE")
    for e in resume.experience:
        para = doc.add_paragraph()
        para.add_run(f"{e.title} — {e.employer}").bold = True
        para.add_run(f"    {e.start} – {e.end}")
        if e.location:
            doc.add_paragraph(e.location)
        for b in e.bullets:
            doc.add_paragraph(b, style="List Bullet")

    if resume.projects:
        heading("AI ENGINEERING PROJECTS")
        for p in resume.projects:
            para = doc.add_paragraph()
            para.add_run(p.name).bold = True
            if p.url and p.shipped:
                para.add_run(f"    {p.url.replace('https://', '')}")
            if p.stack:
                doc.add_paragraph(p.stack)
            for b in p.bullets:
                doc.add_paragraph(b, style="List Bullet")

    if resume.certifications:
        heading("CERTIFICATIONS")
        for c in resume.certifications:
            doc.add_paragraph(c, style="List Bullet")
    if resume.education:
        heading("EDUCATION")
        for ed in resume.education:
            doc.add_paragraph(ed)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _rt(text: str, bold: bool = False) -> list[dict[str, Any]]:
    chunks = [text[i : i + 1900] for i in range(0, len(text), 1900)] or [""]
    return [{"type": "text", "text": {"content": c}, "annotations": {"bold": bold}} for c in chunks]


def markdown_to_blocks(markdown: str) -> list[dict[str, Any]]:
    """A small, predictable Markdown subset -> Notion blocks."""
    blocks: list[dict[str, Any]] = []
    for line in markdown.splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        if s.startswith("### "):
            blocks.append({"type": "heading_3", "heading_3": {"rich_text": _rt(s[4:])}})
        elif s.startswith("## "):
            blocks.append({"type": "heading_2", "heading_2": {"rich_text": _rt(s[3:])}})
        elif s.startswith("# "):
            blocks.append({"type": "heading_1", "heading_1": {"rich_text": _rt(s[2:])}})
        elif s.startswith("- "):
            text = s[2:].replace("**", "")
            blocks.append(
                {"type": "bulleted_list_item", "bulleted_list_item": {"rich_text": _rt(text)}}
            )
        else:
            blocks.append({"type": "paragraph", "paragraph": {"rich_text": _rt(s)}})
    return blocks


def ats_summary_lines(resume: TailoredResume, min_ats: float) -> list[str]:
    a = resume.ats
    lines = []
    if a is None:
        return ["ATS estimate: not computed"]
    state = "READY" if resume.ready(min_ats) else "NOT READY"
    lines.append(
        f"ESTIMATED ATS {a.total:.0f}/100 (minimum {min_ats:.0f}) · validation "
        f"{'passed' if resume.is_final else 'FAILED'} · {state}"
    )
    lines.append(
        "Components: "
        + ", ".join(f"{k.replace('_', ' ')} {v:.0f}" for k, v in a.components.items())
    )
    if a.missing_required:
        lines.append(
            "Required skills the resume cannot truthfully show: " + ", ".join(a.missing_required)
        )
    if a.missing_preferred:
        lines.append("Preferred skills not shown: " + ", ".join(a.missing_preferred))
    if resume.validation:
        for issue in resume.validation.blocking:
            lines.append(f"BLOCKING: {issue.detail}")
    lines.append(
        "This is our estimate, not any employer's real ATS score. Below the minimum, the "
        "system never adds claims to raise it - the job waits for Leela's decision."
    )
    return lines
