"""v1.3: real profile invariants, ATS gate, rendering, community sources
(Telegram, job blogs), link safety, government jobs, the daily digest."""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pytest

from career_agent.ats import ATSEstimator
from career_agent.discovery import safety
from career_agent.discovery.daily import (
    DailyResult,
    build_digest,
    prepare_resumes,
    send_email,
    since,
)
from career_agent.discovery.govt import GovtSink, lead_from_telegram, run_govt, scan_page
from career_agent.discovery.http import HttpError, robots_allows
from career_agent.discovery.notion_sink import (
    MemorySink,
    NotionTrackerSink,
    _rt_links,
    page_id_of,
)
from career_agent.discovery.pipeline import (
    BoardConfig,
    DiscoveryConfig,
    GovtWatch,
    run_discovery,
)
from career_agent.discovery.sources import RawPosting, fetch_jobsite, fetch_telegram
from career_agent.discovery.telegram import (
    channel_trust,
    classify,
    fields,
    first_external_link,
    guess_company_role,
    parse_channel,
    parse_count,
)
from career_agent.discovery.webjobs import detail_links, parse_detail
from career_agent.models import JobPosting
from career_agent.profile_store import load_profile
from career_agent.render import (
    markdown_to_blocks,
    resume_filename,
    to_docx,
    to_markdown,
)
from career_agent.scoring import Scorer
from career_agent.tailor import ResumeTailor

ROOT = Path(__file__).resolve().parents[1]
WEB = ROOT / "tests" / "fixtures" / "web"
TODAY = date(2026, 9, 17)


def web(name: str) -> str:
    return (WEB / name).read_text(encoding="utf-8")


@pytest.fixture(scope="module")
def real():
    return load_profile(ROOT / "profile" / "master_profile.yaml")


# -- the real master profile ---------------------------------------------------
class TestRealProfile:
    def test_unconfirmed_items_are_not_in_the_profile(self, real):
        corpus = real.claim_corpus()
        assert "internal developer tool" not in corpus
        assert "cloud digital leader" not in corpus
        assert not any("google cloud" in c.lower() for c in real.certifications)

    def test_removed_projects_are_absent(self, real):
        names = " ".join(p.name.lower() for p in real.projects)
        assert "muraqib" not in names and "career agent" not in names
        assert {p.name.split(" ")[0] for p in real.projects} == {"mutagent", "mirage"}

    def test_microservices_is_printable_but_gaps_are_not(self, real):
        assert real.skill("microservices").printable
        for gap in ("kubernetes", "langchain", "pytorch", "vector databases", "kafka", "react"):
            assert not real.skill(gap).printable, gap

    def test_unpublished_project_has_no_link(self, real):
        mirage = next(p for p in real.projects if p.name.startswith("mirage"))
        assert mirage.url == "" and mirage.shipped is False

    def test_every_real_tailored_resume_passes_the_gate(self, real):
        for role in ("GenAI Engineer", "Java Backend Developer", "Software Engineer II"):
            job = JobPosting(
                company="Acme",
                role=role,
                min_years=3,
                required_skills=["java", "rest api", "microservices", "python", "rag"],
                description="Build services and LLM features.",
            )
            resume = ResumeTailor(real).tailor(job, Scorer(real).score(job))
            assert resume.is_final, [i.detail for i in resume.validation.issues]


# -- tailoring + ATS -------------------------------------------------------------
def _job(**kw):
    base = {
        "company": "Acme AI",
        "role": "GenAI Engineer",
        "location": "Hyderabad",
        "min_years": 3,
        "required_skills": ["python", "llm", "rag", "mcp", "java", "rest api"],
        "preferred_skills": ["docker"],
        "description": "Build agentic RAG applications with LLM tool calling and evaluation.",
    }
    base.update(kw)
    return JobPosting(**base)


def test_headline_and_summary_follow_the_role_family(real):
    t = ResumeTailor(real)
    ai = t.tailor(_job(), Scorer(real).score(_job()))
    be_job = _job(role="Java Backend Developer")
    be = t.tailor(be_job, Scorer(real).score(be_job))
    assert ai.headline == real.headlines["ai"]
    assert be.headline == real.headlines["backend"]
    assert ai.summary.startswith(real.summary_openers["ai"])
    assert be.summary.startswith(real.summary_openers["backend"])
    assert ai.summary.count("4.5+ years") == 1


def test_ats_is_high_only_when_the_truth_supports_it(real):
    good = _job()
    r = ResumeTailor(real).tailor(good, Scorer(real).score(good))
    assert r.ats and r.ats.total >= 90 and r.ready(90)
    gap = _job(required_skills=["kubernetes", "pytorch", "langchain", "java"], min_years=8)
    g = ResumeTailor(real).tailor(gap, Scorer(real).score(gap))
    assert g.ats.total < 90 and not g.ready(90)
    assert set(g.ats.missing_required) == {"kubernetes", "pytorch", "langchain"}
    printed = {s.lower() for items in g.skills.values() for s in items}
    assert "kubernetes" not in printed and "langchain" not in printed


def test_ats_components_and_edges(real):
    est = ATSEstimator(real)
    job = _job(required_skills=[], preferred_skills=[], min_years=0, description="")
    r = ResumeTailor(real).tailor(job, Scorer(real).score(job))
    a = est.estimate(job, r)
    assert a.components["required_skills"] == pytest.approx(24.5)
    assert a.components["experience"] == pytest.approx(12.0)
    assert any("No named required skills" in n for n in a.notes)
    phd = _job(mandatory_qualifications=["PhD in Computer Science"], min_years=6)
    assert est.estimate(phd, r).components["education"] == 0
    assert est.estimate(phd, r).components["experience"] == pytest.approx(3.75, abs=0.06)
    assert est.estimate(_job(min_years=5.5), r).components["experience"] == pytest.approx(9.0)
    assert est.estimate(_job(min_years=9), r).components["experience"] == 0
    assert est._title("Marketing Lead", "x") == pytest.approx(4.5)
    assert est._title("Java Developer", "AI only") == pytest.approx(9.0)


def test_render_markdown_docx_and_blocks(real):
    job = _job()
    r = ResumeTailor(real).tailor(job, Scorer(real).score(job))
    md = to_markdown(r, real)
    assert "github.com/leelakrishna288/mutagent" in md
    assert "## PROFESSIONAL EXPERIENCE" in md
    data = to_docx(r, real)
    from docx import Document

    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "NAGA LEELA KRISHNA" in text and "Accenture" in text
    assert not Document(io.BytesIO(data)).tables  # single column, ATS-safe
    blocks = markdown_to_blocks(md + "\n# H1\nplain")
    kinds = {b["type"] for b in blocks}
    assert {"heading_1", "heading_2", "heading_3", "bulleted_list_item", "paragraph"} <= kinds
    assert resume_filename(r).startswith("Acme_AI_GenAI_Engineer_") and resume_filename(r).endswith(
        ".docx"
    )


# -- link safety -------------------------------------------------------------------
@pytest.mark.parametrize(
    ("url", "text", "verdict"),
    [
        ("https://job-boards.greenhouse.io/acme/jobs/1", "", safety.OK),
        ("https://careers.example.com/jobs/1", "", safety.OK),
        ("https://www.cdac.in/x.pdf", "", safety.OK),
        ("https://bit.ly/abc", "", safety.CAUTION),
        ("https://www.offcampusjobs4u.com/post", "", safety.CAUTION),
        ("https://example.net/apply", "", safety.CAUTION),
        ("https://wa.me/9199", "", safety.SUSPICIOUS),
        ("http://example.com/apply", "", safety.SUSPICIOUS),
        ("https://jobs.example.xyz/x", "", safety.SUSPICIOUS),
        ("https://xn--80ak6aa92e.com/", "", safety.SUSPICIOUS),
        ("", "", safety.SUSPICIOUS),
        ("https://1.2.3.4/job", "", safety.DANGEROUS),
        ("https://careers.example.com/app.apk", "", safety.DANGEROUS),
        ("https://careers.example.com/j", "Pay Rs 499 registration fee", safety.DANGEROUS),
        (
            "https://careers.example.com/j",
            "Accenture Exam Help Done Successfully",
            safety.DANGEROUS,
        ),
        ("https://careers.example.com/j", "Launch Offer: lifetime access", safety.SUSPICIOUS),
    ],
)
def test_link_safety(url, text, verdict):
    assert safety.assess(url, text).verdict == verdict


class RepClient:
    def __init__(self, sb=None, vt=None, fail=False):
        self.sb, self.vt, self.fail = sb, vt, fail
        self.calls = []

    def request(self, method, url, body=None, headers=None):
        self.calls.append(url)
        if self.fail:
            raise HttpError(500, url)
        if "safebrowsing" in url:
            return self.sb
        return self.vt


def test_reputation_checks_are_optional_and_honest():
    r = safety.assess("https://careers.example.com/j", "", safety.ReputationChecker(RepClient()))
    assert "Safe Browsing not configured" in r.checks and "VirusTotal not configured" in r.checks
    hit = safety.ReputationChecker(
        RepClient(sb={"matches": [{"threatType": "SOCIAL_ENGINEERING"}]}), "k", ""
    )
    assert safety.assess("https://careers.example.com/j", "", hit).verdict == safety.DANGEROUS
    vt = {"data": {"attributes": {"last_analysis_stats": {"malicious": 1, "suspicious": 0}}}}
    r2 = safety.assess(
        "https://careers.example.com/j",
        "",
        safety.ReputationChecker(RepClient(sb={}, vt=vt), "k", "v"),
    )
    assert r2.verdict == safety.SUSPICIOUS and "Safe Browsing clean" in r2.checks
    vt_bad = {"data": {"attributes": {"last_analysis_stats": {"malicious": 3}}}}
    r3 = safety.assess(
        "https://careers.example.com/j", "", safety.ReputationChecker(RepClient(vt=vt_bad), "", "v")
    )
    assert r3.verdict == safety.DANGEROUS
    failed = safety.ReputationChecker(RepClient(fail=True), "k", "v")
    r4 = safety.assess("https://careers.example.com/j", "", failed)
    assert r4.verdict == safety.OK and any("failed" in c or "no report" in c for c in r4.checks)
    # plain-http links are never sent for lookup
    rc = RepClient()
    safety.ReputationChecker(rc, "k", "v").check("http://x", safety.SafetyResult())
    assert rc.calls == []


# -- Telegram --------------------------------------------------------------------
def test_parse_public_channel_preview():
    ch = parse_channel("offcampusjobsindia_IT", web("tg_offcampusjobsindia_IT.html"))
    assert ch.has_preview and len(ch.posts) == 20
    assert ch.subscribers == 99_600 and ch.kind == "subscribers"
    assert ch.title
    assert all(p.url.startswith("https://t.me/offcampusjobsindia_IT/") for p in ch.posts)
    assert any("SIEMENS" in p.text for p in ch.posts)
    assert channel_trust(ch, TODAY, 5000)[0] == "HIGH"


def test_group_without_preview_is_reported_not_guessed():
    ch = parse_channel("offcampusjobs_4u", web("tg_offcampusjobs_4u.html"))
    assert not ch.has_preview and ch.posts == []
    assert ch.subscribers and ch.kind == "members"
    assert channel_trust(ch, TODAY, 5000)[0] == "LOW"


def test_small_or_stale_channels_get_low_trust():
    ch = parse_channel("experiencedjobs", web("tg_experiencedjobs.html"))
    assert channel_trust(ch, TODAY, 5000)[0] == "LOW"
    assert channel_trust(ch, date(2026, 12, 1), 10)[0] == "LOW"
    assert channel_trust(ch, TODAY, 100)[0] == "MEDIUM"


def test_parse_count():
    assert parse_count("99.6K") == 99_600
    assert parse_count("1.2M") == 1_200_000
    assert parse_count("672") == 672
    assert parse_count("n/a") is None


@pytest.mark.parametrize(
    ("text", "kind"),
    [
        (
            "Acme is hiring Java Backend Engineer\nExperience: 3-6 years\nApply: https://x",
            "private",
        ),
        ("Hiring Software Engineer\nBatch: 2025\nFreshers", "skip"),
        ("Medpace Hiring\nRole: Entry Level Software Engineer", "skip"),
        ("Principal Java Architect, 12+ years of experience", "skip"),
        ("Sales executive wanted, java knowledge a plus", "skip"),
        ("Graphic designer needed", "skip"),
        ("", "skip"),
        ("BEL Recruitment 2026: Senior Engineer (Computer Science), B.E/B.Tech CSE", "govt"),
        ("SBI Recruitment 2026 for IT officer", "skip"),
        ("Police constable recruitment by state government", "skip"),
        ("NIT Recruitment 2026 Junior Engineer, B.Tech Civil", "skip"),
    ],
)
def test_classify_posts(text, kind):
    assert classify(text)[0] == kind


def test_fields_and_company_guess():
    text = "🚨 SIEMENS IS HIRING 🔥\n💼 Role: Software Development Engineer\n📍 Location: Bangalore, India"
    f = fields(text)
    assert f["role"] == "Software Development Engineer" and f["location"] == "Bangalore, India"
    assert guess_company_role("Company Name: Acme\nPosition: Backend Dev") == (
        "Acme",
        "Backend Dev",
    )
    c, r = guess_company_role("Zeta is hiring for Java Developer")
    assert c == "Zeta" and "Java Developer" in r


class HtmlClient:
    """get_text/request fake keyed by URL prefix."""

    def __init__(self, pages, json_routes=None, today=TODAY):
        self.pages = pages
        self.json = json_routes or {}
        self.today = today
        self.fetched = []
        self.calls = []

    def get_text(self, url, headers=None):
        self.fetched.append(url)
        for prefix, value in self.pages.items():
            if url.startswith(prefix):
                if isinstance(value, Exception):
                    raise value
                return value
        raise HttpError(404, url)

    def request(self, method, url, body=None, headers=None):
        self.calls.append((method, url, body))
        for prefix, value in self.json.items():
            if url.startswith(prefix):
                return value(body) if callable(value) else value
        raise HttpError(404, url)

    def _sleep(self, s):
        pass


def test_fetch_telegram_keeps_only_safe_relevant_leads():
    post = """
    <div class="tgme_channel_info_header_title"><span>Jobs</span></div>
    <span class="counter_value">20K</span> <span class="counter_type">subscribers</span>
    <div class="tgme_widget_message js-widget_message" data-post="chan/1">
      <div class="tgme_widget_message_text js-message_text">Company: Acme<br/>Role: Java Backend Engineer<br/>Experience: 3-6 years<br/>Location: Hyderabad<br/><a href="https://job-boards.greenhouse.io/acme/jobs/9">Apply</a></div>
      <a class="tgme_widget_message_date"><time datetime="2026-09-17T05:00:00+00:00"></time></a>
    </div>
    <div class="tgme_widget_message js-widget_message" data-post="chan/2">
      <div class="tgme_widget_message_text js-message_text">Company: Scam<br/>Role: Java Developer<br/>Pay Rs 999 registration fee <a href="https://scam.example/x">link</a></div>
      <time datetime="2026-09-17T05:00:00+00:00"></time>
    </div>
    <div class="tgme_widget_message js-widget_message" data-post="chan/3">
      <div class="tgme_widget_message_text js-message_text">Java developer, no link</div>
      <time datetime="2026-09-17T05:00:00+00:00"></time>
    </div>
    <div class="tgme_widget_message js-widget_message" data-post="chan/4">
      <div class="tgme_widget_message_text js-message_text">ECIL Recruitment 2026 Project Engineer (CSE) <a href="https://www.ecil.co.in/jobs/Advt_12_2026.pdf">pdf</a></div>
      <time datetime="2026-09-16T05:00:00+00:00"></time>
    </div>
    <div class="tgme_widget_message js-widget_message" data-post="chan/5">
      <div class="tgme_widget_message_text js-message_text">Old Java Backend Engineer post <a href="https://careers.x.com/1">a</a></div>
      <time datetime="2026-08-01T05:00:00+00:00"></time>
    </div>"""
    from career_agent.discovery.sources import Board

    out = fetch_telegram(HtmlClient({"https://t.me/s/chan": post}), Board("telegram", "chan", "C"))
    assert [o.category for o in out] == ["private", "govt"]
    lead = out[0]
    assert lead.company == "Acme" and lead.title == "Java Backend Engineer"
    assert lead.url == "https://job-boards.greenhouse.io/acme/jobs/9"
    assert lead.source_url == "https://t.me/chan/1" and lead.safety_verdict == safety.OK
    assert lead.source_trust.startswith("HIGH")
    assert out[1].url.endswith("Advt_12_2026.pdf")


def test_fetch_telegram_errors():
    from career_agent.discovery.sources import Board

    with pytest.raises(ValueError, match="no public preview"):
        fetch_telegram(
            HtmlClient({"https://t.me/s/g": web("tg_offcampusjobs_4u.html")}),
            Board("telegram", "g", "G"),
        )

    class JsonOnly:
        def request(self, *a, **k):
            return None

    with pytest.raises(ValueError, match="HTML-capable"):
        fetch_telegram(JsonOnly(), Board("telegram", "g", "G"))
    with pytest.raises(ValueError, match="HTML-capable"):
        fetch_jobsite(JsonOnly(), Board("jobsite", "https://x.test/jobs", "X"))


def test_first_external_link_skips_telegram_links():
    ch = parse_channel("x", web("tg_offcampusjobsindia_IT.html"))
    links = [first_external_link(p) for p in ch.posts]
    assert all("t.me/" not in link for link in links if link)


# -- job blog (JSON-LD) ------------------------------------------------------------
def test_jobsite_listing_and_detail():
    links = detail_links("https://www.offcampusjobsindia.com/category/it-jobs", web("ocji_it.html"))
    assert len(links) >= 5 and all("/jobs/" in u for u in links)
    job = parse_detail(
        "https://www.offcampusjobsindia.com/jobs/accenture-custom-software-engineer-off-campus-hyderabad",
        web("ocji_detail.html"),
    )
    assert job and job.company == "Accenture" and job.title == "Custom Software Engineer"
    assert job.apply_url.startswith("https://www.accenture.com/")
    assert job.country == "IN" and job.posted.startswith("2026-08-19")
    assert parse_detail("https://x.test/jobs/abc", "<html>no data</html>") is None


def test_fetch_jobsite_respects_robots_and_uses_employer_link():
    from career_agent.discovery.sources import Board

    base = "https://www.offcampusjobsindia.com"
    client = HtmlClient(
        {
            f"{base}/robots.txt": "User-Agent: *\nAllow: /\nDisallow: /admin\n",
            f"{base}/category/it-jobs": web("ocji_it.html"),
            f"{base}/jobs/accenture": web("ocji_detail.html"),
            f"{base}/jobs/": "<html></html>",
        }
    )
    out = fetch_jobsite(client, Board("jobsite", f"{base}/category/it-jobs", "ocji"))
    assert out and all(o.source_url.startswith(base) for o in out)
    assert all(not o.url.startswith(base) for o in out)  # employer links, not the blog
    assert out[0].url.startswith("https://www.accenture.com/") and out[0].company == "Accenture"
    blocked = HtmlClient({f"{base}/robots.txt": "User-agent: *\nDisallow: /\n"})
    with pytest.raises(ValueError, match="robots"):
        fetch_jobsite(blocked, Board("jobsite", f"{base}/category/it-jobs", "ocji"))


def test_robots_allows_edge_cases():
    assert robots_allows(HtmlClient({}), "https://none.test/x") is True  # 404 robots
    forbidden = HtmlClient({"https://f.test/robots.txt": HttpError(403, "u")})
    assert robots_allows(forbidden, "https://f.test/x") is False

    class Boom(HtmlClient):
        def get_text(self, url, headers=None):
            raise OSError("down")

    assert robots_allows(Boom({}), "https://b.test/x") is True


# -- pipeline with community sources ----------------------------------------------------
def test_pipeline_routes_govt_leads_and_marks_community_rows_for_review(profile):
    post = """
    <span class="counter_value">50K</span><span class="counter_type">subscribers</span>
    <div class="tgme_widget_message js-widget_message" data-post="c/1">
      <div class="tgme_widget_message_text">Company: Acme<br/>Role: Java Backend Engineer<br/>Experience: 4+ years<br/>Location: Hyderabad<br/>Java, REST API, SQL, Python, LLM, RAG, MCP<br/><a href="https://careers.acme.com/j/1">Apply</a></div>
      <time datetime="2026-09-17T05:00:00+00:00"></time></div>
    <div class="tgme_widget_message js-widget_message" data-post="c/2">
      <div class="tgme_widget_message_text">C-DAC Recruitment 2026 Project Engineer computer science <a href="https://www.cdac.in/a.pdf">pdf</a></div>
      <time datetime="2026-09-17T05:00:00+00:00"></time></div>"""
    cfg = DiscoveryConfig(
        boards=[BoardConfig(platform="telegram", token="c", company="C")],
        include_titles=["nothing matches this"],
        exclude_titles=["principal"],
        target_locations=["hyderabad"],
    )
    sink = MemorySink()
    rep = run_discovery(cfg, profile, HtmlClient({"https://t.me/s/c": post}), sink, today=TODAY)
    assert len(rep.govt_leads) == 1
    assert len(sink.records) == 1
    rec = sink.records[0]
    assert rec.analysis.decision.value in ("MANUAL_REVIEW", "DO_NOT_APPLY")
    assert "Telegram" in rec.notes and "t.me/c/1" in rec.notes
    assert rep.records and rep.records[0][0] is rec


# -- government ----------------------------------------------------------------
GOV_PAGE = """
<a href="/docs/Advt_07_2026_Project_Engineer_CSE.pdf">Advt 07/2026: Project Engineer (Computer Science)</a>
<a href="/docs/result_2026.pdf">Result of Project Engineer 2026</a>
<a href="/docs/Advt_2019.pdf">Advt 2019 Scientist computer science</a>
<a href="/docs/Advt_08_2026_Civil.pdf">Advt 08/2026 Junior Engineer Civil</a>
<a href="mailto:hr@x.gov.in">mail</a>
<a href="https://bank.example/ibps-2026">IBPS IT officer 2026</a>
<a href="/careers/walk-in-2026-software-developer">Walk-in 2026 Software Developer</a>
"""


def test_scan_official_page():
    leads = scan_page("ECIL", "https://www.ecil.co.in/jobs.php", GOV_PAGE, False, TODAY)
    urls = [lead.url for lead in leads]
    assert urls == [
        "https://www.ecil.co.in/docs/Advt_07_2026_Project_Engineer_CSE.pdf",
        "https://www.ecil.co.in/careers/walk-in-2026-software-developer",
    ]
    assert leads[0].verification == "PARTIALLY VERIFIED" and "not been read" in leads[0].notes
    cs = scan_page(
        "C-DAC",
        "https://www.cdac.in/x",
        '<a href="/r/Recruitment_2026_Admin.pdf">Recruitment 2026</a>',
        True,
        TODAY,
    )
    assert len(cs) == 1


def govt_notion(rows, created):
    def query(body):
        return {"results": rows, "has_more": False}

    def create(body):
        created.append(body)
        return {"url": "https://notion.so/g1"}

    return {
        "https://api.notion.com/v1/data_sources/": query,
        "https://api.notion.com/v1/pages": create,
    }


def gov_row(post, url, status="CHECK ELIGIBILITY", last="2026-09-29"):
    t = lambda v: {"type": "rich_text", "rich_text": [{"plain_text": v}]}  # noqa: E731
    return {
        "url": "https://notion.so/row",
        "properties": {
            "Post": {"type": "title", "title": [{"plain_text": post}]},
            "Organisation": t("BEL"),
            "Notification URL": {"type": "url", "url": url},
            "Notes": t("see https://seen.example/x"),
            "Status": {"select": {"name": status}},
            "Last Date": {"date": {"start": last}},
            "Fit": {"select": {"name": "Good fit"}},
        },
    }


def test_run_govt_dedups_writes_and_lists_deadlines():
    created = []
    rows = [
        gov_row("BEL SE", "https://www.ecil.co.in/docs/Advt_07_2026_Project_Engineer_CSE.pdf"),
        gov_row("Old", "https://o.example", status="EXPIRED"),
        gov_row("Past", "https://p.example", last="2026-09-01"),
    ]
    client = HtmlClient(
        {
            "https://www.ecil.co.in/robots.txt": "",
            "https://www.ecil.co.in/jobs.php": GOV_PAGE,
            "https://down.example": HttpError(503, "x"),
            "https://blocked.example/robots.txt": "User-agent: *\nDisallow: /\n",
        },
        govt_notion(rows, created),
    )
    sink = GovtSink(client, "tok", "collection://g", sleep=lambda s: None)
    tg = RawPosting(
        company="Unknown (see https://t.me/c/9)",
        title="NIELIT Scientist B (CS)",
        location="",
        url="https://www.nielit.gov.in/a.pdf",
        external_id="tg-c/9",
        platform="Telegram",
        description="NIELIT Recruitment 2026 Scientist B computer science",
        source_url="https://t.me/c/9",
        source_name="Telegram C",
        source_trust="HIGH (x)",
        safety="Link safety OK",
    )
    scam = RawPosting(
        company="X",
        title="Y",
        location="",
        url="https://x.example",
        external_id="s",
        platform="Telegram",
        description="pay rs 500 registration fee",
    )
    watch = [
        GovtWatch(org="ECIL", url="https://www.ecil.co.in/jobs.php"),
        GovtWatch(org="Down", url="https://down.example/j"),
        GovtWatch(org="Blocked", url="https://blocked.example/j"),
    ]
    rep = run_govt(watch, client, [tg, scam], sink, today=TODAY)
    assert rep.pages_ok == 1 and len(rep.pages_failed) == 2
    assert [lead.url for lead in rep.new] == [
        "https://www.ecil.co.in/careers/walk-in-2026-software-developer",
        "https://www.nielit.gov.in/a.pdf",
    ]
    assert len(created) == 2
    props = created[1]["properties"]
    assert props["Source"]["select"]["name"] == "Telegram"
    assert props["Verification"]["select"]["name"] == "UNVERIFIED"
    assert props["Status"]["select"]["name"] == "NEW"
    assert "date of birth" in props["Leela Must Check"]["rich_text"][0]["text"]["content"]
    assert [d["post"] for d in rep.open_deadlines] == ["BEL SE"]
    assert lead_from_telegram(tg, TODAY).organisation.startswith("NIELIT")
    with pytest.raises(ValueError):
        GovtSink(client, "", "x")


def test_run_govt_without_sink_and_caps():
    client = HtmlClient(
        {"https://www.ecil.co.in/robots.txt": "", "https://www.ecil.co.in/jobs.php": GOV_PAGE}
    )
    rep = run_govt(
        [GovtWatch(org="ECIL", url="https://www.ecil.co.in/jobs.php")],
        client,
        [],
        None,
        today=TODAY,
        max_new=1,
    )
    assert len(rep.new) == 1 and rep.skipped >= 1


# -- Notion helpers -------------------------------------------------------------
def test_page_id_and_links():
    assert (
        page_id_of("https://www.notion.so/Acme-3debe82cb5588160904fcaf9085ed447")
        == "3debe82cb5588160904fcaf9085ed447"
    )
    assert page_id_of("3debe82c-b558-8160-904f-caf9085ed447") == "3debe82cb5588160904fcaf9085ed447"
    assert page_id_of("not-an-id") == "not-an-id"
    rt = _rt_links("apply at https://x.example/a, then relax")["rich_text"]
    assert rt[1]["text"]["link"]["url"] == "https://x.example/a"
    assert rt[0]["text"]["content"] == "apply at " and rt[2]["text"]["content"].startswith(",")
    assert _rt_links("")["rich_text"][0]["text"]["content"] == ""


class NotionFake:
    def __init__(self):
        self.calls = []
        self.uploads = []

    def request(self, method, url, body=None, headers=None):
        self.calls.append((method, url, body))
        if url.endswith("/file_uploads"):
            return {"id": "fu-1"}
        if "/query" in url:
            return {
                "results": [
                    {
                        "url": "https://notion.so/r1",
                        "properties": {
                            "Company": {"type": "rich_text", "rich_text": [{"plain_text": "Acme"}]},
                            "Role": {"type": "rich_text", "rich_text": [{"plain_text": "SE"}]},
                            "Job URL": {"type": "url", "url": "https://j"},
                            "Match Score": {"number": 81},
                            "Status": {"select": {"name": "MATCHED"}},
                            "Company Verdict": {"select": None},
                        },
                    }
                ]
            }
        if url.endswith("/pages"):
            return {"id": "p-new", "url": "https://notion.so/p-new"}
        return {}

    def post_multipart(self, url, field, filename, content, content_type, headers=None):
        self.uploads.append((url, field, filename, len(content)))
        return {"status": "uploaded"}


def test_tracker_sink_resume_helpers():
    fake = NotionFake()
    sink = NotionTrackerSink(fake, "t", "ds", sleep=lambda s: None)
    assert sink.upload_file("a.docx", b"123", "x/y") == "fu-1"
    assert fake.uploads[0][0].endswith("/file_uploads/fu-1/send")
    sink.update_properties("https://notion.so/Row-3debe82cb5588160904fcaf9085ed447", {"x": 1})
    assert fake.calls[-1][1].endswith("/pages/3debe82cb5588160904fcaf9085ed447")
    sink.append_blocks("p1", [{"type": "divider", "divider": {}}] * 95)
    assert sum(1 for c in fake.calls if c[1].endswith("/blocks/p1/children")) == 2
    rows = sink.pending_rows("2026-09-01")
    assert rows[0]["company"] == "Acme" and rows[0]["score"] == 81 and rows[0]["verdict"] == ""
    long_md = "# T\n" + "\n".join(f"- line {i} https://x.example/{i}" for i in range(130))
    sink.write_report("parent", "Daily", long_md)
    assert any(c[1].endswith("/blocks/p-new/children") for c in fake.calls)

    class NoMultipart:
        def request(self, *a, **k):
            return {"id": "fu"}

    with pytest.raises(RuntimeError):
        NotionTrackerSink(NoMultipart(), "t", "ds", sleep=lambda s: None).upload_file("a", b"", "x")

    class NoId:
        def request(self, *a, **k):
            return {}

    with pytest.raises(RuntimeError):
        NotionTrackerSink(NoId(), "t", "ds", sleep=lambda s: None).upload_file("a", b"", "x")


# -- daily run -----------------------------------------------------------------
LONG_JD = (
    "Requirements\n- 3+ years of experience building backend services in Java\n"
    "- REST APIs, SQL, Python, LLM APIs, RAG, MCP, microservices\n"
    + "You will design and operate services with our platform team. "
    * 10
)


def gh_board(jd=LONG_JD, title="Software Engineer, Java (GenAI)"):
    return {
        "jobs": [
            {
                "id": 7,
                "title": title,
                "company_name": "Acme AI",
                "location": {"name": "Hyderabad, India"},
                "absolute_url": "https://job-boards.greenhouse.io/acme/jobs/7",
                "content": jd,
                "first_published": "2026-09-15T00:00:00Z",
            },
            {
                "id": 8,
                "title": "Backend Engineer",
                "company_name": "Acme AI",
                "location": {"name": "Hyderabad, India"},
                "absolute_url": "https://job-boards.greenhouse.io/acme/jobs/8",
                "content": "Requirements\n- 3+ years of experience in Java\n- Kubernetes, PyTorch, LangChain, Kafka, Spark, Scala\n"
                + "Operate data systems. " * 20,
                "first_published": "2026-09-15T00:00:00Z",
            },
        ]
    }


def test_daily_resumes_digest_and_notion_attachment(real):
    cfg = DiscoveryConfig(
        boards=[BoardConfig(platform="greenhouse", token="acme", company="Acme AI")],
        include_titles=["engineer"],
        target_locations=["hyderabad"],
    )
    client = HtmlClient({}, {"https://boards-api.greenhouse.io/v1/boards/acme/": gh_board()})
    sink = MemorySink()
    rep = run_discovery(cfg, real, client, sink, today=TODAY)
    # MemorySink returns no page id, so resumes are made but not attached
    result = DailyResult()
    prepare_resumes(rep, real, cfg, None, result)
    assert result.resumes, rep.recorded
    ready = [r for r in result.resumes if r.ready]
    assert ready and ready[0].docx[:2] == b"PK"
    assert ready[0].auto_submit_eligible == (ready[0].score >= 80)

    fake = NotionFake()
    tracker = NotionTrackerSink(fake, "t", "ds", sleep=lambda s: None)
    rep2 = run_discovery(
        cfg,
        real,
        HtmlClient({}, {"https://boards-api.greenhouse.io/v1/boards/acme/": gh_board()}),
        MemorySink(),
        today=TODAY,
    )
    rep2.records = [(rec, "3debe82cb5588160904fcaf9085ed447") for rec, _ in rep2.records]
    result2 = DailyResult()
    prepare_resumes(rep2, real, cfg, tracker, result2)
    patches = [c for c in fake.calls if c[0] == "PATCH" and "/pages/" in c[1]]
    assert patches and "ATS Estimate" in patches[0][2]["properties"]
    appended = [c for c in fake.calls if c[1].endswith("/children")]
    kinds = [b["type"] for c in appended for b in c[2]["children"]]
    assert "callout" in kinds and "file" in kinds and "heading_2" in kinds
    assert fake.uploads

    digest = build_digest("2026-09-17", rep, result, None)
    assert "Apply now" in digest and "Acme AI" in digest
    assert "Government tracker not configured" in digest
    assert "never auto-applied" in digest


def test_daily_resume_upload_failure_falls_back_to_text(real):
    class NoUpload(NotionFake):
        def post_multipart(self, *a, **k):
            raise HttpError(400, "u", "bad")

    cfg = DiscoveryConfig(
        boards=[BoardConfig(platform="greenhouse", token="acme", company="Acme AI")],
        include_titles=["engineer"],
    )
    rep = run_discovery(
        cfg,
        real,
        HtmlClient({}, {"https://boards-api.greenhouse.io/v1/boards/acme/": gh_board()}),
        MemorySink(),
        today=TODAY,
    )
    rep.records = [(rec, "pid") for rec, _ in rep.records]
    fake = NoUpload()
    result = DailyResult()
    prepare_resumes(
        rep, real, cfg, NotionTrackerSink(fake, "t", "ds", sleep=lambda s: None), result
    )
    texts = [
        rt["text"]["content"]
        for c in fake.calls
        if c[1].endswith("/children")
        for b in c[2]["children"]
        if b["type"] == "paragraph"
        for rt in b["paragraph"]["rich_text"]
    ]
    assert any("DOCX upload failed" in t for t in texts)
    assert not result.errors


def test_digest_sections_with_govt_and_pending(real):
    from career_agent.discovery.daily import ResumeOutcome
    from career_agent.discovery.govt import GovtLead, GovtReport
    from career_agent.discovery.pipeline import RunReport

    rep = RunReport(
        run_date="2026-09-17", boards_ok=1, boards_failed=["X (telegram:g): no preview"]
    )
    res = DailyResult(
        resumes=[
            ResumeOutcome("A", "SE", "https://a", "TARGET", 85, 95, True, "a.docx", [], True),
            ResumeOutcome(
                "B", "SE", "https://b", "PRACTICE", 72, 80, False, "b.docx", ["kafka"], False
            ),
        ],
        pending=[
            {
                "purpose": "TARGET",
                "company": "C",
                "role": "R",
                "score": 77.0,
                "status": "MATCHED",
                "verdict": "",
                "url": "https://c",
            }
        ],
    )
    gov = GovtReport(
        pages_ok=3,
        pages_failed=["BEL: SSL"],
        new=[
            GovtLead(
                "Post",
                "ECIL",
                "https://e",
                "Official site",
                "PARTIALLY VERIFIED",
                "n",
                "2026-09-17",
            )
        ],
        open_deadlines=[
            {
                "post": "BEL SE",
                "last": "2026-09-29",
                "fit": "Good fit",
                "status": "CHECK ELIGIBILITY",
                "url": "",
                "page": "https://notion.so/x",
            }
        ],
    )
    d = build_digest("2026-09-17", rep, res, gov)
    assert "meets auto-submit bar" in d and "missing kafka" in d
    assert "Still waiting" in d and "NEW · ECIL" in d and "Deadline 2026-09-29" in d
    assert "Official pages not reachable today: 1" in d and "Source failed" in d
    empty = build_digest("2026-09-17", rep, DailyResult(), GovtReport())
    assert "None today." in empty and "No new government" in empty
    assert since(TODAY, 14) == "2026-09-03"


def test_send_email_builds_a_message_with_attachments():
    sent = {}

    class FakeSMTP:
        def __init__(self, host, port, context=None, timeout=None):
            sent["host"] = (host, port)

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def login(self, user, password):
            sent["login"] = (user, password)

        def send_message(self, msg):
            sent["msg"] = msg

    send_email("s", "body", "me@x", "pw", "me@x", [("r.docx", b"PK..")], smtp_factory=FakeSMTP)
    msg = sent["msg"]
    assert sent["host"] == ("smtp.gmail.com", 465) and sent["login"] == ("me@x", "pw")
    assert msg["Subject"] == "s" and msg["To"] == "me@x"
    assert [p.get_filename() for p in msg.iter_attachments()] == ["r.docx"]


def test_pipeline_skips_fresher_only_postings(profile):
    board = {
        "jobs": [
            {
                "id": 1,
                "title": "Software Engineer",
                "location": {"name": "Hyderabad"},
                "absolute_url": "https://job-boards.greenhouse.io/a/jobs/1",
                "content": "New grad role for the 2026 batch. Java, SQL.",
                "first_published": "2026-09-10T00:00:00Z",
            }
        ]
    }
    cfg = DiscoveryConfig(
        boards=[BoardConfig(platform="greenhouse", token="a", company="A")],
        include_titles=["software engineer"],
    )
    client = HtmlClient({}, {"https://boards-api.greenhouse.io/v1/boards/a/": board})
    rep = run_discovery(cfg, profile, client, MemorySink(), today=TODAY)
    assert rep.fresher_only == 1 and not rep.recorded


# -- CLI -----------------------------------------------------------------------
def _cli_config(tmp_path, notion_block=""):
    cfg = tmp_path / "d.yaml"
    cfg.write_text(
        "boards:\n  - { platform: greenhouse, token: acme, company: Acme AI }\n"
        "include_titles: [engineer]\ntarget_locations: [hyderabad]\n"
        "govt_watch:\n  - { org: ECIL, url: 'https://www.ecil.co.in/jobs.php' }\n" + notion_block,
        encoding="utf-8",
    )
    return cfg


class CliClient(HtmlClient):
    instance = None

    def __init__(self, *a, **k):
        routes = {
            "https://boards-api.greenhouse.io/v1/boards/acme/": gh_board(),
            "https://api.notion.com/v1/data_sources/": {"results": [], "has_more": False},
            "https://api.notion.com/v1/pages": {"id": "p1", "url": "https://notion.so/p1"},
            "https://api.notion.com/v1/file_uploads": {"id": "fu"},
            "https://api.notion.com/v1/blocks/": {},
        }
        super().__init__(
            {"https://www.ecil.co.in/robots.txt": "", "https://www.ecil.co.in/jobs.php": GOV_PAGE},
            routes,
            today=None,
        )
        CliClient.instance = self

    def post_multipart(self, *a, **k):
        return {}


def test_cli_daily_dry_run_and_notion(tmp_path, monkeypatch):
    from typer.testing import CliRunner

    from career_agent import cli
    from career_agent.discovery import daily as daily_mod
    from career_agent.discovery import http as http_mod

    monkeypatch.setattr(http_mod, "UrllibJsonClient", CliClient)
    runner = CliRunner()
    prof = str(ROOT / "profile" / "master_profile.yaml")
    out = tmp_path / "out"
    res = runner.invoke(
        cli.app,
        [
            "daily",
            "-c",
            str(_cli_config(tmp_path)),
            "-p",
            prof,
            "--out",
            str(out),
            "--save-resumes",
            "--public-summary",
            str(tmp_path / "pub.md"),
            "--email",
        ],
        env={"GMAIL_USER": "", "GMAIL_APP_PASSWORD": ""},
    )
    assert res.exit_code == 0, res.output
    assert "Email skipped" in res.output
    assert list(out.glob("digest_*.md")) and list(out.glob("*.docx"))
    pub = (tmp_path / "pub.md").read_text()
    assert "Acme" not in pub and "Resumes prepared" in pub

    sent = {}
    monkeypatch.setattr(daily_mod, "send_email", lambda *a, **k: sent.setdefault("args", a))
    monkeypatch.setattr(cli, "load_profile", load_profile)
    res2 = runner.invoke(
        cli.app,
        [
            "daily",
            "-c",
            str(
                _cli_config(tmp_path, "notion:\n  report_page_id: rp\n  govt_data_source_id: gds\n")
            ),
            "-p",
            prof,
            "--out",
            str(out),
            "--notion",
            "--email",
        ],
        env={
            "NOTION_TOKEN": "t",
            "NOTION_DATA_SOURCE_ID": "ds",
            "GMAIL_USER": "me@x",
            "GMAIL_APP_PASSWORD": "pw",
        },
    )
    assert res2.exit_code == 0, res2.output
    assert "Digest emailed" in res2.output and sent["args"][2] == "me@x"
    urls = [c[1] for c in CliClient.instance.calls]
    assert any("/data_sources/gds/query" in u for u in urls)

    res3 = runner.invoke(
        cli.app,
        ["daily", "-c", str(_cli_config(tmp_path)), "-p", prof, "--notion"],
        env={"NOTION_TOKEN": "", "NOTION_DATA_SOURCE_ID": ""},
    )
    assert res3.exit_code == 2


# -- HTTP client: text pages and multipart ---------------------------------------
class _Resp:
    def __init__(self, data, charset="utf-8"):
        self._data = data
        self.headers = self

    def get_content_charset(self):
        return "utf-8"

    def read(self, n=-1):
        return self._data

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def test_get_text_retries_and_errors(monkeypatch):
    import urllib.error

    from career_agent.discovery import http as http_mod

    calls = {"n": 0}

    def fake(req, timeout):
        calls["n"] += 1
        if calls["n"] == 1:
            raise urllib.error.HTTPError(req.full_url, 503, "x", {}, io.BytesIO(b""))
        if calls["n"] == 2:
            raise urllib.error.URLError("reset")
        if calls["n"] == 3:
            return _Resp("héllo".encode())
        raise urllib.error.HTTPError(req.full_url, 403, "no", {}, io.BytesIO(b""))

    monkeypatch.setattr(http_mod.urllib.request, "urlopen", fake)
    c = http_mod.UrllibJsonClient(sleep=lambda s: None, retries=2)
    assert c.get_text("https://x.test/") == "héllo"
    with pytest.raises(HttpError):
        c.get_text("https://x.test/")
    with pytest.raises(ValueError):
        c.get_text("http://x.test/")

    def always_down(req, timeout):
        raise urllib.error.URLError("down")

    monkeypatch.setattr(http_mod.urllib.request, "urlopen", always_down)
    with pytest.raises(urllib.error.URLError):
        http_mod.UrllibJsonClient(sleep=lambda s: None, retries=1).get_text("https://x.test/")


def test_post_multipart(monkeypatch):
    import urllib.error

    from career_agent.discovery import http as http_mod

    seen = {}

    def ok(req, timeout):
        seen["ct"] = req.headers["Content-type"]
        seen["body"] = req.data
        return _Resp(b'{"status": "uploaded"}')

    monkeypatch.setattr(http_mod.urllib.request, "urlopen", ok)
    c = http_mod.UrllibJsonClient(sleep=lambda s: None)
    out = c.post_multipart("https://api.test/send", "file", 'a"b.docx', b"DATA", "x/y", {"A": "1"})
    assert out == {"status": "uploaded"}
    assert seen["ct"].startswith("multipart/form-data; boundary=")
    assert b'filename="ab.docx"' in seen["body"] and b"DATA" in seen["body"]
    with pytest.raises(ValueError):
        c.post_multipart("http://api.test/send", "file", "a", b"", "x/y")

    def bad(req, timeout):
        raise urllib.error.HTTPError(req.full_url, 400, "bad", {}, io.BytesIO(b"nope"))

    monkeypatch.setattr(http_mod.urllib.request, "urlopen", bad)
    with pytest.raises(HttpError) as err:
        c.post_multipart("https://api.test/send", "file", "a", b"", "x/y")
    assert err.value.body == "nope"
